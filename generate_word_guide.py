import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_word_guide():
    doc = Document()

    # Page setup - 0.75 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Base styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    # Read markdown content
    md_path = 'VIVA_AND_TEAM_STUDY_GUIDE.md'
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    # Title & Metadata Banner
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("CFM56-5B4 Turbofan Digital Twin")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(14, 116, 144) # Teal / Cyan

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(12)
    run_sub = subtitle_p.add_run("Team Master Study & Technical Viva Guide | Data Warehousing & Mining (Sem 5)")
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Metadata callout box
    meta_table = doc.add_table(rows=1, cols=1)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_cell = meta_table.cell(0, 0)
    meta_cell.width = Inches(7.0)
    set_cell_background(meta_cell, "F1F5F9")
    set_cell_margins(meta_cell, top=140, bottom=140, left=200, right=200)
    mp = meta_cell.paragraphs[0]
    mp.paragraph_format.space_after = Pt(2)
    m_run = mp.add_run("Airframe Context: Airbus A320-214 | Powerplant: CFM56-5B4/P Dual-Spool Turbofan\n"
                       "Dataset: NASA C-MAPSS FD001 (100 Turbofans, 21 Sensors, Run-to-Failure)\n"
                       "Standards Compliance: ATA 100/2200 Spec | FAA Form 8130-3 / EASA Form 1 Release Directive")
    m_run.font.size = Pt(9.5)
    m_run.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Parsing markdown lines into structured Word paragraphs, tables, and callouts
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    for line in lines:
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                # Render code box
                c_table = doc.add_table(rows=1, cols=1)
                c_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                c_cell = c_table.cell(0, 0)
                c_cell.width = Inches(7.0)
                set_cell_background(c_cell, "0F172A") # Dark slate
                set_cell_margins(c_cell, top=120, bottom=120, left=180, right=180)
                cp = c_cell.paragraphs[0]
                cp.paragraph_format.space_after = Pt(0)
                code_text = "".join(code_lines)
                c_run = cp.add_run(code_text.strip())
                c_run.font.name = 'Consolas'
                c_run.font.size = Pt(9)
                c_run.font.color.rgb = RGBColor(226, 232, 240)
                code_lines = []
                doc.add_paragraph().paragraph_format.space_after = Pt(4)
            else:
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Handle Markdown Tables
        if "|" in stripped and not stripped.startswith(">"):
            if not in_table:
                in_table = True
                table_lines = [stripped]
            else:
                table_lines.append(stripped)
            continue
        else:
            if in_table:
                in_table = False
                # Parse and render table
                rows_data = []
                for t_line in table_lines:
                    if "---" in t_line:
                        continue
                    cells = [c.strip() for c in t_line.split("|")[1:-1]]
                    if cells:
                        rows_data.append(cells)
                if rows_data:
                    num_cols = max(len(r) for r in rows_data)
                    word_table = doc.add_table(rows=len(rows_data), cols=num_cols)
                    word_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for r_idx, row in enumerate(rows_data):
                        is_header = (r_idx == 0)
                        for c_idx, cell_val in enumerate(row):
                            if c_idx < num_cols:
                                w_cell = word_table.cell(r_idx, c_idx)
                                if is_header:
                                    set_cell_background(w_cell, "0284C7") # Sky blue header
                                else:
                                    set_cell_background(w_cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
                                set_cell_margins(w_cell, top=80, bottom=80, left=100, right=100)
                                p = w_cell.paragraphs[0]
                                p.paragraph_format.space_after = Pt(0)
                                run = p.add_run(cell_val.replace("**", ""))
                                run.font.size = Pt(9.5 if not is_header else 10)
                                run.font.bold = is_header
                                if is_header:
                                    run.font.color.rgb = RGBColor(255, 255, 255)
                                else:
                                    run.font.color.rgb = RGBColor(30, 41, 59)
                    doc.add_paragraph().paragraph_format.space_after = Pt(4)
                table_lines = []

        # Headings
        if stripped.startswith("# "):
            continue # Already handled in title
        elif stripped.startswith("## "):
            h_text = stripped[3:].strip()
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            r = h.add_run(h_text)
            r.font.name = 'Arial'
            r.font.size = Pt(15)
            r.font.bold = True
            r.font.color.rgb = RGBColor(2, 132, 199) # Bright Blue
            continue
        elif stripped.startswith("### "):
            h_text = stripped[4:].strip()
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            h.paragraph_format.keep_with_next = True
            r = h.add_run(h_text)
            r.font.name = 'Arial'
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42) # Slate Dark
            continue

        # Blockquote (Callout / Viva Pitch)
        if stripped.startswith("> "):
            q_text = stripped[2:].strip().replace("**", "").replace("*", "")
            q_table = doc.add_table(rows=1, cols=1)
            q_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            q_cell = q_table.cell(0, 0)
            q_cell.width = Inches(7.0)
            set_cell_background(q_cell, "EFF6FF") # Light blue-50
            set_cell_margins(q_cell, top=100, bottom=100, left=180, right=180)
            qp = q_cell.paragraphs[0]
            qp.paragraph_format.space_after = Pt(0)
            q_run = qp.add_run(q_text)
            q_run.font.size = Pt(10)
            q_run.font.italic = True
            q_run.font.color.rgb = RGBColor(29, 78, 216)
            doc.add_paragraph().paragraph_format.space_after = Pt(3)
            continue

        # Bullet points
        if stripped.startswith("* ") or stripped.startswith("- "):
            bp = doc.add_paragraph(style='List Bullet')
            bp.paragraph_format.space_after = Pt(2)
            text_body = stripped[2:].strip()
            # Simple bold parsing
            parts = text_body.split("**")
            for i, part in enumerate(parts):
                run = bp.add_run(part)
                if i % 2 == 1:
                    run.font.bold = True
            continue

        # Numbered list
        if len(stripped) > 2 and stripped[0].isdigit() and stripped[1] in [".", ")"]:
            np = doc.add_paragraph(style='List Number')
            np.paragraph_format.space_after = Pt(2)
            text_body = stripped[2:].strip()
            parts = text_body.split("**")
            for i, part in enumerate(parts):
                run = np.add_run(part)
                if i % 2 == 1:
                    run.font.bold = True
            continue

        # Normal text paragraph
        if stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            parts = stripped.split("**")
            for i, part in enumerate(parts):
                run = p.add_run(part)
                if i % 2 == 1:
                    run.font.bold = True

    output_file = "CFM56_Digital_Twin_Team_Viva_Guide.docx"
    doc.save(output_file)
    print(f"Successfully generated Word document: {output_file}")

if __name__ == '__main__':
    create_word_guide()
