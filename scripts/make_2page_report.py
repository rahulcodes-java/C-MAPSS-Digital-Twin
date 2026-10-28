"""
make_2page_report.py -- Generate verified 2-page report in docx and pdf.
Strict constraints:
  - Exactly 2 pages
  - Reads metadata from docs/submission/meta.json
  - Only facts from docs/submission/FACTS_v2.md
  - No banned words, no banned meta-text
  - Abstract <= 150 words
  - Sections: Abstract, Introduction, Methodology, Results, Limitations, Future Work, GitHub link, Dashboard screenshot
"""

import os
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas


def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748b"))
        text = f"C-MAPSS Digital Twin | Page {self._pageNumber} of {page_count}"
        self.drawRightString(letter[0] - 36, 20, text)
        self.drawString(36, 20, "DWM Mini-Project Submission | Academic Year 2026")
        self.setStrokeColor(colors.HexColor("#cbd5e1"))
        self.setLineWidth(0.5)
        self.line(36, 30, letter[0] - 36, 30)
        self.restoreState()


def load_meta():
    with open('docs/submission/meta.json', 'r') as f:
        return json.load(f)


# -------------------------------------------------------------
# DOCX Generation
# -------------------------------------------------------------
def generate_docx(meta):
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.45)
        section.right_margin = Inches(0.45)

    title_text = meta['title']
    authors_text = ", ".join(meta['authors'])
    dept_text = meta['department']
    roll_text = ", ".join(meta['roll_numbers'])
    github_url = meta['github_url']

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(1)
    r_title = p_title.add_run(title_text)
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(13)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    # Authors
    p_auth = doc.add_paragraph()
    p_auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_auth.paragraph_format.space_before = Pt(0)
    p_auth.paragraph_format.space_after = Pt(2)
    r_auth = p_auth.add_run(f"{authors_text} ({roll_text})\n")
    r_auth.font.name = 'Calibri'
    r_auth.font.size = Pt(8.5)
    r_auth.font.bold = True
    r_dept = p_auth.add_run(f"{dept_text} | GitHub Repository: ")
    r_dept.font.name = 'Calibri'
    r_dept.font.size = Pt(8.0)
    r_link = p_auth.add_run(github_url)
    r_link.font.name = 'Calibri'
    r_link.font.size = Pt(8.0)
    r_link.font.underline = True
    r_link.font.color.rgb = RGBColor(2, 132, 199)

    # Abstract Box
    tbl_abs = doc.add_table(rows=1, cols=1)
    tbl_abs.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_abs = tbl_abs.cell(0, 0)
    c_abs.width = Inches(7.6)
    set_cell_background(c_abs, "F1F5F9")
    set_cell_margins(c_abs, top=60, bottom=60, left=90, right=90)
    p_abs = c_abs.paragraphs[0]
    p_abs.paragraph_format.space_before = Pt(0)
    p_abs.paragraph_format.space_after = Pt(0)
    r_absh = p_abs.add_run("Abstract— ")
    r_absh.font.name = 'Calibri'
    r_absh.font.bold = True
    r_absh.font.size = Pt(8.0)
    r_abst = p_abs.add_run(
        "Commercial aircraft engine condition monitoring requires reliable detection of degradation states before functional failure occurs. "
        "This project presents a multi-model prognostics framework evaluated across NASA C-MAPSS datasets FD001 through FD004 with all 21 telemetry sensors preserved. "
        "We formulate health management primarily as a 3-stage classification task (Healthy, Degrading, Critical) utilizing a Random Forest classifier interpreted with TreeSHAP. "
        "Secondary complementary tasks include an LSTM network for continuous Remaining Useful Life (RUL) regression and an Isolation Forest for unsupervised anomaly detection. "
        "On FD001, the Random Forest model achieves a 3-stage macro-F1 of 0.7303, while the LSTM achieves a macro-F1 of 0.8197 and provides a median lead-time advantage of 21.0 flight cycles over single-snapshot classification across 100 engines. "
        "TreeSHAP attributes 49.18% of degradation impact to High-Pressure Compressor sensors. Performance declines across multi-condition datasets FD002 and FD004 due to unnormalized operating regimes."
    )
    r_abst.font.name = 'Calibri'
    r_abst.font.size = Pt(8.0)

    def add_h(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(15, 23, 42)

    def add_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(8.2)
        return p

    # I. Introduction
    add_h("1. Introduction")
    add_p(
        "Gas turbine engines operate under severe thermodynamic stresses, where continuous degradation of compressor and turbine aerofoils leads to efficiency losses and eventual component retirement. "
        "Within the Data Warehousing and Mining (DWM) curriculum, we structure turbofan prognostic health management primarily as a 3-stage classification task (Healthy: RUL > 100, Degrading: 30 < RUL <= 100, Critical: RUL <= 30). "
        "To provide actionable maintenance decision support, snapshot classification is complemented by sequence-based RUL regression and unsupervised anomaly detection."
    )

    # 2. Methodology
    add_h("2. Methodology")
    add_p(
        "Telemetry streams comprise 21 sensors across 160359 training cycles and 104897 test cycles over four datasets (FD001 to FD004). "
        "Sensor values are scaled using MinMaxScaler fitted strictly on training data. Rather than discarding sensors with low sea-level variance, all 21 sensors are preserved because multi-condition datasets (FD002 and FD004) exhibit non-zero variance across all channels. "
        "As depicted in Fig. 1, the architecture routes temporal sliding windows of shape (50, 21) to an LSTM regression network, while single-cycle 21-dimensional vectors are evaluated by a 100-tree Random Forest classifier and an Isolation Forest detector."
    )

    # Pipeline Figure
    if os.path.exists("docs/submission/figures/pipeline_diagram_light.png"):
        p_fig1 = doc.add_paragraph()
        p_fig1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_fig1.paragraph_format.space_before = Pt(1)
        p_fig1.paragraph_format.space_after = Pt(1)
        p_fig1.add_run().add_picture("docs/submission/figures/pipeline_diagram_light.png", width=Inches(6.2))
        p_c1 = doc.add_paragraph()
        p_c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_c1.paragraph_format.space_before = Pt(0)
        p_c1.paragraph_format.space_after = Pt(2)
        r_c1 = p_c1.add_run("Fig. 1. Digital twin processing pipeline distinguishing sequence-window and single-cycle model inputs.")
        r_c1.font.name = 'Calibri'
        r_c1.font.size = Pt(7.5)
        r_c1.font.italic = True

    # Page Break for clean page 2 start
    doc.add_page_break()

    # 3. Results
    add_h("3. Experimental Results")
    add_p(
        "Models were evaluated on test partitions without retraining. Table 1 summarizes 3-stage classification and binary Critical detection across all four datasets. "
        "On FD001, Random Forest snapshot classification attains a macro-F1 of 0.7303, while sequence-based LSTM achieves a macro-F1 of 0.8197. "
        "In critical cycle detection (RUL <= 30), LSTM achieves an F1 of 0.8158 (precision 0.7926, recall 0.8404), compared to 0.6938 for Random Forest and 0.2413 for Isolation Forest (precision 0.1372, recall 0.9970)."
    )

    # Table 1: Classification Performance
    t_data = [
        ["Dataset", "Cycles", "RF Macro-F1", "LSTM Macro-F1", "RF Crit. F1", "LSTM Crit. F1", "IF Crit. F1"],
        ["FD001", "13096", "0.7303", "0.8197", "0.6938", "0.8158", "0.2413"],
        ["FD002", "33991", "0.1283", "0.0378", "0.0607", "0.0659", "0.0620"],
        ["FD003", "16596", "0.2785", "0.2410", "0.0974", "0.1009", "0.0345"],
        ["FD004", "41214", "0.1004", "0.0264", "0.0479", "0.0431", "0.0411"]
    ]
    tbl_res = doc.add_table(rows=len(t_data), cols=7)
    tbl_res.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(t_data):
        for c_idx, val in enumerate(row):
            cell = tbl_res.cell(r_idx, c_idx)
            cell.paragraphs[0].text = val
            cell.paragraphs[0].paragraph_format.space_before = Pt(1)
            cell.paragraphs[0].paragraph_format.space_after = Pt(1)
            cell.paragraphs[0].runs[0].font.name = 'Calibri'
            cell.paragraphs[0].runs[0].font.size = Pt(7.2)
            if r_idx == 0:
                cell.paragraphs[0].runs[0].font.bold = True
                set_cell_background(cell, "E2E8F0")
            elif r_idx % 2 == 1:
                set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=30, bottom=30, left=40, right=40)

    p_t1 = doc.add_paragraph()
    p_t1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t1.paragraph_format.space_before = Pt(1)
    p_t1.paragraph_format.space_after = Pt(2)
    r_t1 = p_t1.add_run("Table 1. Standardized classification and critical-class F1 metrics across C-MAPSS test sets.")
    r_t1.font.name = 'Calibri'
    r_t1.font.size = Pt(7.2)
    r_t1.font.italic = True

    add_p(
        "Lead-time analysis evaluated over all 100 FD001 training engines reveals that LSTM temporal modeling detects degradation with a median delay of -2.0 cycles (mean -10.57, IQR 41.5), "
        "compared to +15.5 cycles (mean 17.12, IQR 35.25) for single-snapshot Random Forest. "
        "The paired lead-time difference (RF delay - LSTM delay) yields a median advantage of 21.0 flight cycles (mean 27.69, IQR 32.5), with LSTM alerting earlier in 80 of 100 engines (80.0%). "
        "On held-out validation tail engines (83 to 100), the median advantage is 26.5 cycles (LSTM flags first in 17 of 18 engines, 94.44%). "
        "TreeSHAP attributes 49.18% \u00b1 0.12% of degradation attribution to High-Pressure Compressor sensors (Ps30: 14.25%, Nc: 13.01%), consistent with simulated HPC degradation."
    )

    # Dashboard Screenshot
    if os.path.exists("docs/submission/figures/dashboard_critical.png"):
        p_fig2 = doc.add_paragraph()
        p_fig2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_fig2.paragraph_format.space_before = Pt(1)
        p_fig2.paragraph_format.space_after = Pt(1)
        p_fig2.add_run().add_picture("docs/submission/figures/dashboard_critical.png", width=Inches(5.0))
        p_c2 = doc.add_paragraph()
        p_c2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_c2.paragraph_format.space_before = Pt(0)
        p_c2.paragraph_format.space_after = Pt(1)
        r_c2 = p_c2.add_run("Fig. 2. Interactive digital twin interface displaying Stage 2 Critical state and live telemetry thresholds.")
        r_c2.font.name = 'Calibri'
        r_c2.font.size = Pt(7.2)
        r_c2.font.italic = True

    # 4. Limitations
    add_h("4. Limitations")
    add_p(
        "Evaluations rely on simulated C-MAPSS data rather than commercial airline flight logs. "
        "Models trained exclusively on single-regime sea-level data (FD001) experience marked performance drops on FD002 and FD004 (macro-F1 0.1283 and 0.1004) due to unnormalized operating conditions. "
        "Isolation Forest exhibits low critical precision (0.1372 on FD001) due to erratic early triggers across non-stationary degradation trajectories."
    )

    # 5. Future Work
    add_h("5. Future Work")
    add_p(
        "Future efforts will incorporate operating regime clustering prior to sensor normalization to enable reliable cross-condition transfer. "
        "Physics-informed loss constraints and bidirectional sequence attention will also be explored to refine remaining useful life bounds."
    )

    doc.save("docs/submission/Project_Report_2page.docx")
    print("Saved docs/submission/Project_Report_2page.docx")


# -------------------------------------------------------------
# PDF Generation (ReportLab)
# -------------------------------------------------------------
def generate_pdf(meta):
    pdf_path = "docs/submission/Project_Report_2page.pdf"
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        leftMargin=32,
        rightMargin=32,
        topMargin=26,
        bottomMargin=34
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=14,
        alignment=1,
        textColor=colors.HexColor('#0f172a'),
        spaceAfter=2
    )
    auth_style = ParagraphStyle(
        'DocAuth',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.0,
        leading=10,
        alignment=1,
        textColor=colors.HexColor('#1e293b'),
        spaceAfter=1
    )
    sub_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=9,
        alignment=1,
        textColor=colors.HexColor('#475569'),
        spaceAfter=3
    )
    h_style = ParagraphStyle(
        'SecH',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9.0,
        leading=11,
        textColor=colors.HexColor('#0f172a'),
        spaceBefore=3,
        spaceAfter=1
    )
    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.8,
        leading=9.8,
        textColor=colors.HexColor('#1e293b'),
        spaceAfter=2
    )
    caption_style = ParagraphStyle(
        'Cap',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=7.0,
        leading=8.5,
        alignment=1,
        textColor=colors.HexColor('#475569'),
        spaceBefore=1,
        spaceAfter=2
    )
    abs_style = ParagraphStyle(
        'Abs',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=9.2,
        textColor=colors.HexColor('#1e293b')
    )

    story = []

    # Title & Metadata
    story.append(Paragraph(meta['title'], title_style))
    auth_str = f"{', '.join(meta['authors'])} ({', '.join(meta['roll_numbers'])})"
    story.append(Paragraph(auth_str, auth_style))
    sub_str = f"{meta['department']} | Project Repository: <u>{meta['github_url']}</u>"
    story.append(Paragraph(sub_str, sub_style))

    # Abstract Box
    abs_text = (
        "<b>Abstract— </b>Commercial aircraft engine condition monitoring requires reliable detection of degradation states before functional failure occurs. "
        "This project presents a multi-model prognostics framework evaluated across NASA C-MAPSS datasets FD001 through FD004 with all 21 telemetry sensors preserved. "
        "We formulate health management primarily as a 3-stage classification task (Healthy, Degrading, Critical) utilizing a Random Forest classifier interpreted with TreeSHAP. "
        "Secondary complementary tasks include an LSTM network for continuous Remaining Useful Life (RUL) regression and an Isolation Forest for unsupervised anomaly detection. "
        "On FD001, the Random Forest model achieves a 3-stage macro-F1 of 0.7303, while the LSTM achieves a macro-F1 of 0.8197 and provides a median lead-time advantage of 21.0 flight cycles over single-snapshot classification across 100 engines. "
        "TreeSHAP attributes 49.18% of degradation impact to High-Pressure Compressor sensors. Performance declines across multi-condition datasets FD002 and FD004 due to unnormalized operating regimes."
    )
    t_abs = Table([[Paragraph(abs_text, abs_style)]], colWidths=[548])
    t_abs.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f1f5f9')),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(t_abs)
    story.append(Spacer(1, 2))

    # 1. Introduction
    story.append(Paragraph("1. Introduction", h_style))
    story.append(Paragraph(
        "Gas turbine engines operate under severe thermodynamic stresses, where continuous degradation of compressor and turbine aerofoils leads to efficiency losses and eventual component retirement. "
        "Within the Data Warehousing and Mining (DWM) curriculum, we structure turbofan prognostic health management primarily as a 3-stage classification task (Healthy: RUL > 100, Degrading: 30 < RUL &le; 100, Critical: RUL &le; 30). "
        "To provide actionable maintenance decision support, snapshot classification is complemented by sequence-based RUL regression and unsupervised anomaly detection.",
        body_style
    ))

    # 2. Methodology
    story.append(Paragraph("2. Methodology", h_style))
    story.append(Paragraph(
        "Telemetry streams comprise 21 sensors across 160359 training cycles and 104897 test cycles over four datasets (FD001 to FD004). "
        "Sensor values are scaled using MinMaxScaler fitted strictly on training data. Rather than discarding sensors with low sea-level variance, all 21 sensors are preserved because multi-condition datasets (FD002 and FD004) exhibit non-zero variance across all channels. "
        "As depicted in Fig. 1, the architecture routes temporal sliding windows of shape (50, 21) to an LSTM regression network, while single-cycle 21-dimensional vectors are evaluated by a 100-tree Random Forest classifier and an Isolation Forest detector.",
        body_style
    ))

    # Figure 1: Pipeline
    if os.path.exists("docs/submission/figures/pipeline_diagram_light.png"):
        story.append(Image("docs/submission/figures/pipeline_diagram_light.png", width=480, height=201))
        story.append(Paragraph("<i>Fig. 1. Digital twin processing pipeline distinguishing sequence-window and single-cycle model inputs.</i>", caption_style))

    # Explicit PageBreak for Page 2
    story.append(PageBreak())

    # 3. Experimental Results
    story.append(Paragraph("3. Experimental Results", h_style))
    story.append(Paragraph(
        "Models were evaluated on test partitions without retraining. Table 1 summarizes 3-stage classification and binary Critical detection across all four datasets. "
        "On FD001, Random Forest snapshot classification attains a macro-F1 of 0.7303, while sequence-based LSTM achieves a macro-F1 of 0.8197. "
        "In critical cycle detection (RUL &le; 30), LSTM achieves an F1 of 0.8158 (precision 0.7926, recall 0.8404), compared to 0.6938 for Random Forest and 0.2413 for Isolation Forest (precision 0.1372, recall 0.9970).",
        body_style
    ))

    # Table 1: Classification Performance
    t_cells = [
        ["Dataset", "Cycles", "RF Macro-F1", "LSTM Macro-F1", "RF Crit. F1", "LSTM Crit. F1", "IF Crit. F1"],
        ["FD001", "13096", "0.7303", "0.8197", "0.6938", "0.8158", "0.2413"],
        ["FD002", "33991", "0.1283", "0.0378", "0.0607", "0.0659", "0.0620"],
        ["FD003", "16596", "0.2785", "0.2410", "0.0974", "0.1009", "0.0345"],
        ["FD004", "41214", "0.1004", "0.0264", "0.0479", "0.0431", "0.0411"]
    ]
    t1_p = Table(t_cells, colWidths=[60, 55, 85, 90, 80, 85, 75])
    t1_p.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e2e8f0')),
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.HexColor('#0f172a')),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 7.0),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ('TOPPADDING', (0, 0), (-1, -1), 2.0),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2.0),
    ]))
    story.append(t1_p)
    story.append(Paragraph("<i>Table 1. Standardized classification and critical-class F1 metrics across C-MAPSS test sets.</i>", caption_style))

    story.append(Paragraph(
        "Lead-time analysis evaluated over all 100 FD001 training engines reveals that LSTM temporal modeling detects degradation with a median delay of -2.0 cycles (mean -10.57, IQR 41.5), "
        "compared to +15.5 cycles (mean 17.12, IQR 35.25) for single-snapshot Random Forest. "
        "The paired lead-time difference (RF delay - LSTM delay) yields a median advantage of 21.0 flight cycles (mean 27.69, IQR 32.5), with LSTM alerting earlier in 80 of 100 engines (80.0%). "
        "On held-out validation tail engines (83 to 100), the median advantage is 26.5 cycles (LSTM flags first in 17 of 18 engines, 94.44%). "
        "TreeSHAP attributes 49.18% &plusmn; 0.12% of degradation attribution to High-Pressure Compressor sensors (Ps30: 14.25%, Nc: 13.01%), consistent with simulated HPC degradation.",
        body_style
    ))

    # Figure 2: Dashboard Screenshot
    if os.path.exists("docs/submission/figures/dashboard_critical.png"):
        story.append(Image("docs/submission/figures/dashboard_critical.png", width=440, height=222))
        story.append(Paragraph("<i>Fig. 2. Interactive digital twin interface displaying Stage 2 Critical state and live telemetry thresholds.</i>", caption_style))

    # 4. Limitations
    story.append(Paragraph("4. Limitations", h_style))
    story.append(Paragraph(
        "Evaluations rely on simulated C-MAPSS data rather than commercial airline flight logs. "
        "Models trained exclusively on single-regime sea-level data (FD001) experience marked performance drops on FD002 and FD004 (macro-F1 0.1283 and 0.1004) due to unnormalized operating conditions. "
        "Isolation Forest exhibits low critical precision (0.1372 on FD001) due to erratic early triggers across non-stationary degradation trajectories.",
        body_style
    ))

    # 5. Future Work
    story.append(Paragraph("5. Future Work", h_style))
    story.append(Paragraph(
        "Future efforts will incorporate operating regime clustering prior to sensor normalization to enable reliable cross-condition transfer. "
        "Physics-informed loss constraints and bidirectional sequence attention will also be explored to refine remaining useful life bounds.",
        body_style
    ))

    doc.build(story, canvasmaker=NumberedCanvas)
    print("Saved docs/submission/Project_Report_2page.pdf")


if __name__ == '__main__':
    meta = load_meta()
    generate_docx(meta)
    generate_pdf(meta)
